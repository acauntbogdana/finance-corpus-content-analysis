#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Нейтральный контент-анализ собранного корпуса (unsupervised learning).
Читает research_corpus/records.csv и локальные PDF. Никаких сетевых запросов
или загрузок не делает.

Методы:
1) TF-IDF (слова и биграммы) — частотность с поправкой на распространённость.
2) NMF — латентные темы без заранее заданных тематических категорий.
3) KMeans — диагностика кластерной структуры (выводится как содержательный
   результат только при silhouette >= SILHOUETTE_THRESHOLD).
4) Частоты терминов и фраз — воспроизводимые словари по каждой теме.
5) Сравнение по годам и журналам с нормировкой на число публикаций.

Словари задают только поисковые признаки, а не предписывают вывод.
Выводы описательные. Таблицы 1–5 формируются в DOCX для каждой из шести тем;
сохраняются CSV с документными матрицами, ключевыми словами NMF и назначенными
кластерами, плюс раздел «Ограничения» с автопроверкой.

Установка:
    python -m pip install pandas scikit-learn python-docx pymupdf

Запуск:
    python analyze_finance_corpus.py
    python analyze_finance_corpus.py --max-features 6000 --topics 6
    python analyze_finance_corpus.py --no-pdf-text
"""
from __future__ import annotations

import argparse
import csv
import json
import logging
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

import pandas as pd
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from sklearn.cluster import KMeans
from sklearn.decomposition import NMF
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics import silhouette_score

ROOT = Path(__file__).resolve().parent
CORPUS = ROOT / "research_corpus"
RECORDS_CSV = CORPUS / "records.csv"
OUT = CORPUS / "analysis"

TOPIC_NAMES = {
    "asset_pricing": "Ценообразование активов",
    "discreteness": "Дискретность и целочисленные модели в финансах",
    "waves": "Волны, циклы и вейвлет-анализ финансовых рынков",
    "portfolio_management": "Портфельное управление",
    "investment": "Инвестиции",
    "trading": "Трейдинг",
}

# Ниже этого silhouette KMeans-разбиение не считается содержательной
# кластерной структурой и выводится только как диагностика.
SILHOUETTE_THRESHOLD = 0.10

# Термины, характерные для разделов благодарностей/финансирования в PDF,
# а не для содержания статей. Используются только для пометки артефактных
# NMF-компонент (не как фильтр по документам).
ARTIFACT_TERMS = {
    "grant", "grants", "funding", "fund", "funds", "foundation", "foundations",
    "support", "supported", "acknowledge", "acknowledges", "acknowledgement",
    "acknowledgements", "acknowledgments", "thank", "thanks", "grateful",
    "national", "natural", "china", "chinese", "ministry", "council",
    "university", "universities", "project", "projects", "program", "programs",
    "education", "contract", "agreement", "award", "awards", "sponsor",
    "sponsored", "number", "numbered", "nsfc", "nsf", "erc", "dfg",
    "supplemental", "supplementary", "material", "materials",
    "accepted", "available", "online", "files", "file",
    "downloaded", "permission", "rights", "copyright", "license",
}

TERMS: dict[str, dict[str, list[str]]] = {
    "asset_pricing": {
        "Модели факторов и риска": ["factor model", "fama french", "risk factor", "risk premium", "beta"],
        "Дисконт-фактор и ядро ценообразования": ["stochastic discount factor", "pricing kernel", "discount factor", "euler equation"],
        "CAPM и равновесное ценообразование": ["capital asset pricing", "capm", "market portfolio"],
        "Аномалии и предсказуемость": ["return predictability", "momentum", "value premium", "anomaly", "anomalies"],
        "Ликвидность и транзакционные издержки": ["liquidity", "illiquidity", "transaction cost", "market impact"],
        "Машинное обучение": ["machine learning", "deep learning", "neural network"],
    },
    "discreteness": {
        "Дискретное время и динамические решения": ["discrete time", "multi-period", "rebalancing", "time step"],
        "Целочисленная оптимизация": ["integer programming", "mixed integer", "integer optimization", "lot size"],
        "Кардинальные ограничения": ["cardinality constraint", "cardinality constrained", "minimum number of assets"],
        "Сценарии и дискретные распределения": ["discrete distribution", "scenario tree", "scenario analysis", "binomial tree"],
        "Метаэвристики": ["genetic algorithm", "particle swarm", "metaheuristic", "simulated annealing"],
        "Дискретизация и вычислительные методы": ["discretization", "finite difference", "numerical scheme"],
    },
    "waves": {
        "Вейвлеты и время-частота": ["wavelet", "time frequency", "multiscale", "multi scale"],
        "Рыночные циклы": ["financial cycle", "business cycle", "market cycle", "long wave"],
        "Волатильность и распространение шоков": ["volatility spillover", "contagion", "shock transmission", "volatility wave"],
        "Технические волновые подходы": ["elliott wave", "technical analysis", "wave theory"],
        "Периодичность и спектральные методы": ["spectral analysis", "frequency domain", "periodicity", "fourier"],
    },
    "portfolio_management": {
        "Оптимизация портфеля": ["portfolio optimization", "portfolio selection", "mean variance", "efficient frontier"],
        "Распределение активов и диверсификация": ["asset allocation", "diversification", "strategic allocation", "tactical allocation"],
        "Ребалансировка и динамическое управление": ["rebalancing", "dynamic portfolio", "dynamic asset allocation", "portfolio turnover"],
        "Риск и устойчивость": ["portfolio risk", "value at risk", "expected shortfall", "tail risk", "robust optimization"],
        "ESG и устойчивые инвестиции": ["esg", "sustainable investment", "climate risk", "green finance"],
        "Оценка эффективности": ["sharpe ratio", "performance evaluation", "risk adjusted return", "alpha"],
    },
    "investment": {
        "Инвестиционные решения и стратегии": ["investment decision", "investment strategy", "investment horizon", "investment behavior"],
        "Активы и распределение капитала": ["asset allocation", "capital allocation", "institutional investor", "retail investor"],
        "Риск и доходность": ["risk return", "investment risk", "expected return", "downside risk"],
        "Пассивные и активные стратегии": ["active management", "passive investing", "index fund", "exchange traded fund"],
        "ESG и климатические факторы": ["esg", "sustainable investment", "climate risk", "green investment"],
        "Поведенческие и информационные факторы": ["investor sentiment", "behavioral finance", "information asymmetry", "attention"],
        "ИИ и автоматизация": ["machine learning", "robo advisor", "artificial intelligence", "algorithmic"],
    },
    "trading": {
        "Алгоритмическая и высокочастотная торговля": ["algorithmic trading", "high frequency trading", "hft", "automated trading"],
        "Стратегии и арбитраж": ["trading strategy", "momentum", "mean reversion", "pairs trading", "arbitrage"],
        "Исполнение и рыночное воздействие": ["trade execution", "slippage", "market impact", "transaction cost"],
        "Микроструктура и ликвидность": ["market microstructure", "bid ask spread", "order book", "liquidity"],
        "Машинное обучение и обучение с подкреплением": ["reinforcement learning", "machine learning", "deep learning", "neural network"],
        "Поведение и сентимент": ["investor sentiment", "news sentiment", "herding", "behavioral"],
        "Цифровые активы": ["cryptocurrency", "bitcoin", "crypto market", "digital asset"],
    },
}

STOPWORDS = set("""
a an the and or but if then of in on at by for to from with without into over under
is are was were be been being this that these those it its we our their they them as
can could may might will would should have has had do does did not no more most less
such also both each other which who whom whose where when how than very between among
using use used based results result study studies paper article model models method
methods approach approaches analysis evidence data new proposed show shows shown
grant grants funding fund funds foundation foundations support supported
acknowledge acknowledges acknowledgement acknowledgements acknowledgments
thank thanks grateful national natural china chinese ministry council
university universities project projects program programs education
contract agreement award awards sponsor sponsored nsfc nsf erc dfg
we our us authors author et al ibid e g i e pp vol no number doi http https www
first second third one two three however moreover furthermore therefore thus
supplemental supplementary material materials accepted available files file
online journal elsevier wiley springer press downloaded permission rights
copyright license licence article published publication article's
""".split())


def _norm_journal(name: str) -> str:
    """Нормализует название журнала для сопоставления allowlist ↔ факт."""
    n = (name or "").strip().lower()
    for prefix in ("the ", "a ", "an "):
        if n.startswith(prefix):
            n = n[len(prefix):]
            break
    return re.sub(r"\s+", " ", n)


def clean_text(text: str) -> str:
    text = (text or "").replace("\x00", " ")
    text = re.sub(r"https?://\S+", " ", text)
    text = re.sub(r"\bdoi\s*:\s*\S+", " ", text, flags=re.I)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def pdf_text(path: Path) -> str:
    try:
        import pymupdf as fitz
        doc = fitz.open(path)
        chunks = []
        for page in doc:
            chunks.append(page.get_text("text"))
        doc.close()
        text = "\n".join(chunks)

        # Отрезаем разделы, не относящиеся к содержанию статьи. Без этого
        # PDF засоряет TF-IDF терминами фондов, грантов и названий статей.
        cutoff_patterns = [
            r"\n\s*Acknowledg(e)?ments?\b",
            r"\n\s*Funding\b",
            r"\n\s*Declaration of (Competing |Conflicting )?Interest",
            r"\n\s*Declaration of Competing Interest",
            r"\n\s*References\b",
            r"\n\s*Bibliography\b",
            r"\n\s*Appendix\b",
        ]
        cut_pos = len(text)
        for pat in cutoff_patterns:
            m = re.search(pat, text, re.I)
            if m and m.start() < cut_pos:
                cut_pos = m.start()
        text = text[:cut_pos]

        # Убираем оставшиеся в основном тексте шаблоны благодарностей
        # и финансирования.
        for pat in (
            r"(?i)we (thank|are grateful to)[^.]*\.(?:\s|$)",
            r"(?i)the authors (thank|acknowledge|are grateful)[^.]*\.(?:\s|$)",
            r"(?i)this (work|research|paper|study) (was|is|has been) (supported|funded)[^.]*\.(?:\s|$)",
            r"(?i)(financial|research) support (from|by)[^.]*\.(?:\s|$)",
            r"(?i)national natural science foundation of china[^.]*\.(?:\s|$)",
            r"(?i)\b(grant|fund|project|contract|award)\s+(no\.?|number)?\s*[\w\-]+",
        ):
            text = re.sub(pat, " ", text)

        # Reject low-quality extraction instead of treating boilerplate as findings.
        useful = len(re.findall(r"[A-Za-z]{3,}", text))
        return clean_text(text) if useful >= 150 else ""
    except Exception as exc:
        logging.warning("Не удалось извлечь текст из %s: %s", path.name, exc)
        return ""


def load_records(use_pdf: bool) -> pd.DataFrame:
    if not RECORDS_CSV.exists():
        raise SystemExit(
            f"Не найден файл {RECORDS_CSV}. Сначала запустите harvest_open_finance_papers.py."
        )
    df = pd.read_csv(RECORDS_CSV, keep_default_na=False)
    if "topic" not in df.columns:
        raise SystemExit("В records.csv отсутствует колонка topic.")
    for col in ("title", "abstract", "keywords", "journal", "authors", "doi", "pdf_path"):
        if col not in df.columns:
            df[col] = ""
    if use_pdf:
        cache: dict[str, str] = {}
        texts = []
        for _, row in df.iterrows():
            rel = str(row.get("pdf_path", "") or "")
            extracted = ""
            if rel:
                path = CORPUS / rel
                if path.exists():
                    if rel not in cache:
                        cache[rel] = pdf_text(path)
                    extracted = cache[rel]
            texts.append(extracted)
        df["pdf_text"] = texts
    else:
        df["pdf_text"] = ""
    df["year"] = pd.to_numeric(df.get("year", 0), errors="coerce").fillna(0).astype(int)
    return df


def document_text(row: Any) -> str:
    # Abstract/keywords/title сохраняются даже если PDF-текст пуст.
    return clean_text(" ".join([
        str(row.get("title", "")), str(row.get("abstract", "")),
        str(row.get("keywords", "")), str(row.get("pdf_text", ""))[:50000],
    ]))


def phrase_count(text: str, phrase: str) -> int:
    # Точное совпадение фраз после нормализации пунктуации; без инференса тональности.
    t = re.sub(r"[^a-z0-9]+", " ", text.lower())
    p = re.sub(r"[^a-z0-9]+", " ", phrase.lower()).strip()
    if not p:
        return 0
    return len(re.findall(r"(?<!\w)" + re.escape(p) + r"(?!\w)", t))


def sentence_candidates(
    records: list[dict[str, Any]], topic_key: str, limit: int = 12
) -> list[tuple[dict[str, Any], str, str]]:
    phrases = [p for group in TERMS[topic_key].values() for p in group]
    anchor = [p for p in phrases if len(p.split()) >= 2]
    rows = []
    definition_patterns = [
        r"\bdefined as\b", r"\brefers to\b", r"\bwe define\b",
        r"\bdefinition of\b", r"\bcan be defined\b", r"\bdenotes\b",
        r"\bis the process of\b", r"\bis a method for\b",
    ]
    for rec in records:
        source = str(rec.get("abstract", "") or "")
        if not source:
            source = str(rec.get("pdf_text", "") or "")[:12000]
        sentences = re.split(r"(?<=[.!?])\s+(?=[A-Z(])", source)
        for sentence in sentences:
            s = sentence.strip()
            if len(s) < 45 or len(s) > 700:
                continue
            hits = sum(phrase_count(s, p) for p in anchor)
            defhit = any(re.search(p, s, re.I) for p in definition_patterns)
            if hits or defhit:
                # Кандидатный скор, а не автоматическое утверждение,
                # что предложение является определением.
                score = 2 * int(defhit) + hits + min(len(s.split()) / 100, 1)
                rows.append((
                    rec, s,
                    f"candidate_score={score:.2f}; definition_pattern={defhit}; "
                    f"topic_phrase_hits={hits}",
                ))
    rows.sort(
        key=lambda x: float(re.search(r"candidate_score=([0-9.]+)", x[2]).group(1)),
        reverse=True,
    )
    unique, out = set(), []
    for rec, sent, note in rows:
        key = re.sub(r"\W+", "", sent.lower())
        if key in unique:
            continue
        unique.add(key)
        out.append((rec, sent, note))
        if len(out) >= limit:
            break
    return out


def build_matrix(texts: list[str], max_features: int):
    n = len(texts)
    if n >= 50:
        min_df_val = 3       # отсекаем редкие и артефактные термины
    elif n >= 8:
        min_df_val = 2
    else:
        min_df_val = 1
    max_df_val = 0.85 if n >= 8 else 1.0

    vectorizer = TfidfVectorizer(
        lowercase=True, strip_accents="unicode", stop_words=list(STOPWORDS),
        ngram_range=(1, 2), min_df=min_df_val, max_df=max_df_val,
        max_features=max_features,
        token_pattern=r"(?u)\b[a-zA-Z][a-zA-Z0-9-]{2,}\b",
    )
    try:
        matrix = vectorizer.fit_transform(texts)
    except ValueError:
        vectorizer = TfidfVectorizer(
            lowercase=True, strip_accents="unicode",
            ngram_range=(1, 2), max_features=max_features,
        )
        matrix = vectorizer.fit_transform(texts or ["empty corpus"])
    return vectorizer, matrix


def unsupervised_topics(texts: list[str], max_features: int, requested_topics: int):
    vectorizer, X = build_matrix(texts, max_features)
    terms = vectorizer.get_feature_names_out()
    n_docs, n_features = X.shape
    n_topics = max(1, min(requested_topics, n_docs, n_features))
    if n_docs < 3 or n_features < 3 or X.nnz == 0:
        return vectorizer, X, [], [0] * n_docs, []

    nmf = NMF(
        n_components=n_topics, init="nndsvda", random_state=17,
        max_iter=500, l1_ratio=0.05,
    )
    nmf.fit_transform(X)
    H = nmf.components_
    topic_rows: list[tuple[int, str, bool]] = []
    for idx, comp in enumerate(H):
        top = comp.argsort()[::-1][:12]
        top_terms = [str(terms[i]) for i in top]
        # Компонент считаем артефактным, если >=3 из 12 ведущих терминов
        # относятся к типовым PDF-шумам (funding, grant, foundation, ...).
        artifact_hits = sum(
            1 for t in top_terms
            if any(a in ARTIFACT_TERMS for a in re.split(r"[\s\-]+", t.lower()))
        )
        is_artifact = artifact_hits >= 3
        topic_rows.append((idx, ", ".join(top_terms), is_artifact))

    # KMeans с консервативным k. Silhouette считается только при валидных условиях.
    k = min(max(2, round(n_docs ** 0.5)), n_docs - 1, 8)
    labels = [0] * n_docs
    silhouette: list[float] = []
    if n_docs >= 4 and X.shape[1] >= 2 and X.nnz:
        km = KMeans(n_clusters=k, random_state=17, n_init=10)
        labels = km.fit_predict(X).tolist()
        try:
            if len(set(labels)) > 1 and len(set(labels)) < n_docs:
                silhouette = [round(float(silhouette_score(X, labels)), 4)]
        except Exception:
            silhouette = []
    return vectorizer, X, topic_rows, labels, silhouette


_XML_ILLEGAL = re.compile(
    # Управляющие символы и прочие недопустимые в XML 1.0 code points.
    "[\x00-\x08\x0B\x0C\x0E-\x1F\x7F-\x9F\uFFFE\uFFFF]"
)


def xml_safe(value: Any) -> str:
    """Убирает символы, недопустимые в DOCX/XML. Заменяет их пробелом."""
    text = str(value) if value is not None else ""
    text = text.replace("\x00", " ")
    text = _XML_ILLEGAL.sub(" ", text)
    return text


def add_table(doc: Document, headers: list[str], rows: list[list[Any]]) -> None:
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = xml_safe(h)
        for run in t.rows[0].cells[i].paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row[:len(headers)]):
            cells[i].text = xml_safe(value)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(8)


def save_topic_report(
    topic_key: str, df: pd.DataFrame, max_features: int, n_topics: int
) -> None:
    topic_name = TOPIC_NAMES[topic_key]
    topic_df = df[df["topic"].astype(str).apply(lambda x: topic_key in x.split(";"))].copy()
    if topic_df.empty:
        logging.warning("Нет записей для темы %s — отчёт не создан", topic_key)
        return
    topic_df["analysis_text"] = topic_df.apply(document_text, axis=1)
    topic_df = topic_df[topic_df["title"].astype(str).str.strip().ne("")].copy()

    texts = topic_df["analysis_text"].fillna("").tolist()
    vectorizer, X, topic_rows, cluster_labels, silhouette = unsupervised_topics(
        texts, max_features, n_topics
    )
    topic_df["cluster_kmeans"] = cluster_labels[:len(topic_df)]
    topic_df["document_length_chars"] = [len(x) for x in texts]

    cats = TERMS[topic_key]
    phrase_totals: Counter[str] = Counter()
    phrase_docs: Counter[str] = Counter()
    journal_totals: dict[str, Counter[str]] = defaultdict(Counter)
    year_totals: dict[int, Counter[str]] = defaultdict(Counter)
    for _, row in topic_df.iterrows():
        text = row["analysis_text"]
        for category, phrases in cats.items():
            count = sum(phrase_count(text, p) for p in phrases)
            phrase_totals[category] += count
            if count:
                phrase_docs[category] += 1
            journal_totals[str(row.get("journal") or "не указано")][category] += count
            year_totals[int(row.get("year", 0))][category] += count

    out_dir = OUT / topic_key
    out_dir.mkdir(parents=True, exist_ok=True)
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10)
    head = doc.add_heading(f"Контент-анализ: {topic_name}", 1)
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(
        f"Период публикаций: {topic_df['year'].min()}–{topic_df['year'].max()}. "
        f"В корпусе: {len(topic_df)} записей. Документ сформирован автоматически."
    )
    doc.add_paragraph(
        "Принцип нейтральности: анализ фиксирует наблюдаемые слова, фразы, публикационные "
        "характеристики и статистические группировки. Он не оценивает авторов и не выводит "
        "причинность из частот. Словари категорий заданы заранее и используются только как "
        "один воспроизводимый слой анализа; NMF и KMeans дают отдельный unsupervised-срез."
    )
    doc.add_heading("Методика и ограничения", 2)
    doc.add_paragraph(
        "Единица анализа — публикация; доступный текст включает заголовок, аннотацию, "
        "ключевые слова и извлечённый текст локального OA PDF, если он был получен. "
        "TF-IDF выделяет термины, относительно характерные для корпуса. NMF формирует "
        "латентные темы, KMeans группирует документы по TF-IDF-векторам. Названия тем "
        "интерпретируются исследователем по верхним терминам; алгоритмы сами по себе не "
        "доказывают смысловую принадлежность. Для PDF-текстов возможны ошибки извлечения; "
        "разделы Acknowledgments/Funding/References отсекаются, но остаточные артефакты "
        "возможны и помечаются в таблице 5."
    )
    doc.add_heading("Результаты", 2)

    # ---------------- Table 1: кандидатные определения ----------------
    doc.add_paragraph(
        "Таблица 1. Кандидатные предложения для ручного анализа определений и трактовок"
    )
    records = topic_df.to_dict("records")
    candidates = sentence_candidates(records, topic_key)
    rows1 = []
    for i, (rec, sentence, note) in enumerate(candidates, 1):
        rows1.append([
            i,
            f"{rec.get('authors', '')} ({rec.get('year', '')})",
            rec.get("title", ""),
            sentence,
            note,
        ])
    if not rows1:
        rows1 = [[
            "—", "—", "—",
            "Кандидатные предложения не найдены в доступных аннотациях/текстах.",
            "Не интерпретировать как отсутствие теоретических определений.",
        ]]
    add_table(
        doc,
        ["№", "Автор(ы), год", "Публикация", "Кандидатное предложение", "Критерий отбора"],
        rows1,
    )
    doc.add_paragraph(
        "Примечание: это кандидаты для проверки, а не автоматически подтверждённые "
        "определения. Отсутствие найденных предложений не означает отсутствия определения "
        "в статье."
    )

    # ---------------- Table 2: операциональный словарь ----------------
    doc.add_paragraph("Таблица 2. Операциональный словарь контент-анализа")
    rows2 = [
        [cat, "; ".join(phrases), phrase_totals[cat], phrase_docs[cat],
         round(phrase_docs[cat] / len(topic_df), 3) if len(topic_df) else 0]
        for cat, phrases in cats.items()
    ]
    add_table(
        doc,
        ["Категория", "Поисковые фразы", "Число упоминаний",
         "Публикаций с ≥1 упоминанием", "Доля публикаций"],
        rows2,
    )

    # ---------------- Table 3: частоты по журналам с нормировкой ----------------
    doc.add_paragraph(
        "Таблица 3. Частотность категорий по журналам: абсолютные совпадения и "
        "нормировка на число публикаций журнала в теме"
    )
    journal_docs = Counter(
        str(row.get("journal") or "не указано") for _, row in topic_df.iterrows()
    )
    top_journals = sorted(
        journal_totals,
        key=lambda j: sum(journal_totals[j].values()),
        reverse=True,
    )[:20]
    rows3 = []
    for journal in top_journals:
        total = sum(journal_totals[journal].values())
        docs = journal_docs.get(journal, 0)
        rows3.append(
            [journal, docs]
            + [journal_totals[journal][cat] for cat in cats]
            + [total, round(total / docs, 2) if docs else 0]
        )
    total_docs = len(topic_df)
    total_hits = sum(phrase_totals.values())
    rows3.append(
        ["ИТОГО", total_docs]
        + [phrase_totals[cat] for cat in cats]
        + [total_hits, round(total_hits / total_docs, 2) if total_docs else 0]
    )
    add_table(
        doc,
        ["Журнал", "Публикаций журнала в теме"]
        + list(cats)
        + ["Всего совпадений", "Совпадений на публикацию"],
        rows3,
    )
    doc.add_paragraph(
        "Примечание: «Публикаций журнала в теме» — число записей этого журнала, "
        "прошедших тематический фильтр; это не общее число статей журнала за период. "
        "Абсолютные числа совпадений зависят от размера журнала, поэтому для сравнения "
        "используется нормировка «Совпадений на публикацию»."
    )

    # ---------------- Table 4: динамика по годам ----------------
    doc.add_paragraph("Таблица 4. Динамика частотности категорий по годам публикации")
    years = sorted(y for y in year_totals if y)
    rows4 = [
        [cat] + [year_totals[y][cat] for y in years] + [phrase_totals[cat]]
        for cat in cats
    ]
    rows4.append(
        ["Число публикаций"]
        + [int((topic_df["year"] == y).sum()) for y in years]
        + [len(topic_df)]
    )
    add_table(doc, ["Категория"] + years + ["Всего"], rows4)

    # ---------------- Table 5: NMF + диагностика KMeans ----------------
    sil_val: float | None = silhouette[0] if silhouette else None
    doc.add_paragraph(
        "Таблица 5. Темы NMF и диагностика кластерной структуры KMeans "
        "(unsupervised learning)"
    )
    rows5: list[list[Any]] = []
    for row in topic_rows:
        idx, words, is_artifact = row
        note = "Термины с наибольшим весом в компоненте"
        if is_artifact:
            note = "⚠ АРТЕФАКТ: термины фондов/благодарностей в PDF, не содержательная тема"
        rows5.append([f"NMF {idx + 1}", words, note])

    if sil_val is None:
        rows5.append([
            "KMeans", "не выполнен",
            "Недостаточно документов или нулевая матрица",
        ])
    elif sil_val < SILHOUETTE_THRESHOLD:
        rows5.append([
            "KMeans", f"Silhouette={sil_val}",
            f"Кластерная структура не выявлена (порог {SILHOUETTE_THRESHOLD}). "
            f"Разбиение не интерпретируется как содержательная группировка.",
        ])
    else:
        counts = Counter(cluster_labels)
        for cluster, count in sorted(counts.items()):
            rows5.append([
                f"KMeans кластер {cluster}",
                f"{count} публикаций",
                "Интерпретировать по top terms и названиям статей; "
                "кластер не является готовой категорией",
            ])
        rows5.append([
            "KMeans диагностика", f"Silhouette={sil_val}",
            "Сравнивать только при одинаковой предобработке и параметрах",
        ])
    if not rows5:
        rows5 = [[
            "Недостаточно данных", "NMF/KMeans не выполнены",
            "Нужно больше документов и непустого текста",
        ]]
    add_table(doc, ["Тип/номер", "Ведущие термины или размер", "Комментарий"], rows5)

    # ---------------- Описательные наблюдения ----------------
    doc.add_heading("Краткие описательные наблюдения", 2)
    top = phrase_totals.most_common(3)
    doc.add_paragraph(xml_safe(
        "Наибольшее число совпадений в словаре: "
        + (", ".join(f"{name} ({count})" for name, count in top) if top else "нет совпадений")
        + ". Это отражает частоту заданных фраз в доступном тексте, но не доказывает "
          "важность, эффективность или научное превосходство соответствующего подхода."
    ))
    doc.add_paragraph(
        f"Средняя длина доступного текста: {topic_df['document_length_chars'].mean():.0f} знаков. "
        f"Публикаций с извлечённой аннотацией: "
        f"{(topic_df['abstract'].astype(str).str.len() > 0).sum()} из {len(topic_df)}."
    )

    # ---------------- Ограничения (автопроверка) ----------------
    doc.add_heading("Ограничения текущего анализа", 2)
    limitations: list[str] = []
    limitations.append(
        f"Корпус по теме включает {len(topic_df)} публикаций. Это число не равно общему "
        f"числу статей по теме в этих журналах: тематический фильтр работает по вхождению "
        f"фраз в название и аннотацию, поэтому часть работ с иной терминологией может быть "
        f"недопредставлена."
    )
    if sil_val is not None:
        if sil_val < SILHOUETTE_THRESHOLD:
            limitations.append(
                f"KMeans не выявил устойчивой кластерной структуры "
                f"(silhouette ≈ {sil_val}, порог {SILHOUETTE_THRESHOLD}). "
                f"Результаты KMeans не следует интерпретировать как содержательную "
                f"группировку подтем."
            )
        else:
            limitations.append(
                f"Silhouette для KMeans ≈ {sil_val}, что выше порога "
                f"{SILHOUETTE_THRESHOLD}, но интерпретация кластеров требует ручной "
                f"верификации по названиям статей."
            )
    artifact_count = sum(1 for row in topic_rows if len(row) >= 3 and row[2])
    if artifact_count:
        limitations.append(
            f"NMF выявил {artifact_count} артефактн(ых) компонент(ы), отражающих термины "
            f"фондов, грантов и благодарностей в PDF-файлах, а не тематическую структуру "
            f"корпуса. Такие компоненты помечены в таблице 5 значком «⚠ АРТЕФАКТ» и не "
            f"используются в содержательных выводах."
        )

    allowlist_path = CORPUS.parent / "journal_allowlist.csv"
    missing_journals: list[str] = []
    if allowlist_path.exists():
        try:
            with allowlist_path.open(encoding="utf-8-sig", newline="") as f:
                expected = {
                    row["journal"].strip()
                    for row in csv.DictReader(f)
                    if row.get("include", "yes").strip().lower() in {"yes", "y", "1", "true"}
                }
            present = {str(j).strip() for j in journal_totals if j and j != "не указано"}
            expected_map = {_norm_journal(j): j for j in expected}
            present_norm = {_norm_journal(j) for j in present}
            missing_journals = sorted(
                expected_map[k] for k in expected_map if k not in present_norm
            )
        except Exception as exc:
            logging.warning("Не удалось проверить allowlist: %s", exc)
    if missing_journals:
        limitations.append(
            "В отчёте отсутствуют журналы из allowlist: "
            + ", ".join(missing_journals[:10])
            + (" и др." if len(missing_journals) > 10 else "")
            + ". Причины — недоступность источника метаданных (например, HTTP 401 от "
              "ScienceDirect API) или отсутствие подходящих публикаций в этом журнале "
              "за период."
        )
    for text in limitations:
        doc.add_paragraph(xml_safe(text), style="List Bullet")

    # ---------------- Список публикаций ----------------
    doc.add_heading("Список публикаций", 2)
    for i, rec in enumerate(
        topic_df.sort_values(["year", "journal", "title"]).to_dict("records"), 1
    ):
        authors = rec.get("authors") or "Авторы не указаны"
        title = rec.get("title") or "Название не указано"
        journal = rec.get("journal") or "Журнал не указан"
        year = rec.get("year") or ""
        doi = rec.get("doi") or ""
        p = doc.add_paragraph(xml_safe(f"{i}. {authors} ({year}). {title}. {journal}."))
        if doi:
            p.add_run(xml_safe(f" DOI: {doi}"))

    report_path = out_dir / f"report_{topic_key}.docx"
    doc.save(report_path)

    # ---------------- CSV для воспроизводимости ----------------
    topic_df.drop(columns=["analysis_text"], errors="ignore").to_csv(
        out_dir / "records_analyzed.csv", index=False, encoding="utf-8-sig"
    )
    terms = vectorizer.get_feature_names_out()
    try:
        dense = X.toarray()
        matrix_df = pd.DataFrame(dense, columns=terms)
        # Если термин "title" случайно попал в словарь, не конфликтуем с колонкой-меткой.
        matrix_df = matrix_df.drop(columns=["title"], errors="ignore")
        matrix_df.insert(0, "_document_title", topic_df["title"].values[:len(matrix_df)])
        matrix_df.to_csv(
            out_dir / "tfidf_document_term.csv", index=False, encoding="utf-8-sig"
        )
    except Exception as exc:
        logging.warning("Не сохранена TF-IDF матрица: %s", exc)

    with (out_dir / "unsupervised_topics.csv").open(
        "w", newline="", encoding="utf-8-sig"
    ) as f:
        w = csv.writer(f)
        w.writerow(["method", "topic_or_cluster", "top_terms_or_size", "artifact"])
        for row in topic_rows:
            if len(row) == 3:
                idx, words, is_art = row
            else:
                idx, words = row[0], row[1]
                is_art = False
            w.writerow(["NMF", idx + 1, words, bool(is_art)])
        for cl, count in sorted(Counter(cluster_labels).items()):
            w.writerow(["KMeans", cl, f"documents={count}", False])
        if sil_val is not None:
            w.writerow(["Diagnostics", "silhouette", sil_val, False])

    logging.info("Создан отчёт %s", report_path)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Локальный нейтральный анализ корпуса публикаций."
    )
    parser.add_argument("--max-features", type=int, default=6000)
    parser.add_argument(
        "--topics", type=int, default=6, help="число NMF-компонент на тему"
    )
    parser.add_argument(
        "--no-pdf-text", action="store_true",
        help="анализировать только метаданные/аннотации",
    )
    args = parser.parse_args()
    logging.basicConfig(
        level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s"
    )
    OUT.mkdir(parents=True, exist_ok=True)
    df = load_records(use_pdf=not args.no_pdf_text)
    if df.empty:
        raise SystemExit("Корпус пуст. Сначала соберите публикации.")
    for topic in TOPIC_NAMES:
        try:
            save_topic_report(topic, df, args.max_features, args.topics)
        except Exception:
            logging.exception("Ошибка анализа темы %s", topic)
    summary = {
        "records": int(len(df)),
        "years": sorted(int(y) for y in df["year"].unique() if y),
        "topics": {
            t: int(df["topic"].astype(str).apply(lambda x: t in x.split(";")).sum())
            for t in TOPIC_NAMES
        },
        "method": "TF-IDF + NMF + KMeans (diagnostic) + phrase counts",
        "silhouette_threshold": SILHOUETTE_THRESHOLD,
        "network_access": False,
    }
    (OUT / "analysis_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    logging.info("Анализ завершён: %s", OUT)


if __name__ == "__main__":
    main()
